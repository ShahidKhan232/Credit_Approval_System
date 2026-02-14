from celery import shared_task
from django.db import transaction
from decimal import Decimal
import openpyxl
from datetime import datetime
import os

from .models import Customer, Loan


@shared_task
def ingest_customer_data(file_path='customer_data.xlsx'):
    """
    Background task to ingest customer data from Excel file
    
    Args:
        file_path: Path to customer_data.xlsx file
    """
    print(f"Starting customer data ingestion from {file_path}")
    
    if not os.path.exists(file_path):
        print(f"Error: File {file_path} not found")
        return {'status': 'error', 'message': f'File {file_path} not found'}
    
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        
        customers_created = 0
        customers_updated = 0
        
        # Skip header row
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:  # Skip empty rows
                continue
            
            customer_id = int(row[0])
            first_name = str(row[1])
            last_name = str(row[2])
            phone_number = str(row[3])
            monthly_salary = Decimal(str(row[4]))
            approved_limit = Decimal(str(row[5]))
            current_debt = Decimal(str(row[6])) if row[6] else Decimal('0')
            
            # Create or update customer
            customer, created = Customer.objects.update_or_create(
                customer_id=customer_id,
                defaults={
                    'first_name': first_name,
                    'last_name': last_name,
                    'phone_number': phone_number,
                    'monthly_salary': monthly_salary,
                    'approved_limit': approved_limit,
                    'current_debt': current_debt,
                    'age': 30  # Default age as it's not in the Excel
                }
            )
            
            if created:
                customers_created += 1
            else:
                customers_updated += 1
        
        result = {
            'status': 'success',
            'customers_created': customers_created,
            'customers_updated': customers_updated
        }
        print(f"Customer data ingestion completed: {result}")
        return result
        
    except Exception as e:
        print(f"Error during customer data ingestion: {str(e)}")
        return {'status': 'error', 'message': str(e)}


@shared_task
def ingest_loan_data(file_path='loan_data.xlsx'):
    """
    Background task to ingest loan data from Excel file
    
    Args:
        file_path: Path to loan_data.xlsx file
    """
    print(f"Starting loan data ingestion from {file_path}")
    
    if not os.path.exists(file_path):
        print(f"Error: File {file_path} not found")
        return {'status': 'error', 'message': f'File {file_path} not found'}
    
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        
        loans_created = 0
        loans_updated = 0
        loans_skipped = 0
        
        # Skip header row
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:  # Skip empty rows
                continue
            
            customer_id = int(row[0])
            loan_id = int(row[1])
            loan_amount = Decimal(str(row[2]))
            tenure = int(row[3])
            interest_rate = Decimal(str(row[4]))
            monthly_repayment = Decimal(str(row[5]))
            emis_paid_on_time = int(row[6])
            
            # Parse dates
            start_date = row[7]
            if isinstance(start_date, str):
                start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
            elif isinstance(start_date, datetime):
                start_date = start_date.date()
            
            end_date = row[8]
            if isinstance(end_date, str):
                end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
            elif isinstance(end_date, datetime):
                end_date = end_date.date()
            
            # Check if customer exists
            try:
                customer = Customer.objects.get(customer_id=customer_id)
            except Customer.DoesNotExist:
                print(f"Customer {customer_id} not found, skipping loan {loan_id}")
                loans_skipped += 1
                continue
            
            # Create or update loan
            loan, created = Loan.objects.update_or_create(
                loan_id=loan_id,
                defaults={
                    'customer': customer,
                    'loan_amount': loan_amount,
                    'tenure': tenure,
                    'interest_rate': interest_rate,
                    'monthly_repayment': monthly_repayment,
                    'emis_paid_on_time': emis_paid_on_time,
                    'start_date': start_date,
                    'end_date': end_date
                }
            )
            
            if created:
                loans_created += 1
            else:
                loans_updated += 1
        
        result = {
            'status': 'success',
            'loans_created': loans_created,
            'loans_updated': loans_updated,
            'loans_skipped': loans_skipped
        }
        print(f"Loan data ingestion completed: {result}")
        return result
        
    except Exception as e:
        print(f"Error during loan data ingestion: {str(e)}")
        return {'status': 'error', 'message': str(e)}


@shared_task
def ingest_all_data():
    """
    Background task to ingest both customer and loan data
    """
    print("Starting full data ingestion")
    
    # Ingest customers first
    customer_result = ingest_customer_data.apply()
    
    # Then ingest loans
    loan_result = ingest_loan_data.apply()
    
    return {
        'customer_ingestion': customer_result.get(),
        'loan_ingestion': loan_result.get()
    }
